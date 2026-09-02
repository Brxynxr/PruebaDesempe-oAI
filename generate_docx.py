import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # Cover / Header
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Academia Lumina AI\nDocumentación Técnica Exhaustiva del Sistema")
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1a, 0x36, 0x5d) # Deep Navy

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run("Asistente Inteligente de Atención al Cliente con RAG, Ciberseguridad y Automatización Omnicanal\n")
    sub_run.font.name = 'Arial'
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = RGBColor(0x4a, 0x55, 0x68)

    # Metadata Box
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        ("Autor:", "Breyner Manga"),
        ("Lema:", "El que sabe sabe"),
        ("Versión:", "2.0 (Producción Hardened)"),
        ("Institución:", "Academia Lumina / Riwi Lingua (Colombia)")
    ]

    for idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_k = row.cells[0]
        cell_v = row.cells[1]
        
        cell_k.width = Inches(2.0)
        cell_v.width = Inches(4.5)

        p_k = cell_k.paragraphs[0]
        p_k.paragraph_format.space_after = Pt(2)
        r_k = p_k.add_run(k)
        r_k.bold = True
        r_k.font.color.rgb = RGBColor(0x1a, 0x36, 0x5d)

        p_v = cell_v.paragraphs[0]
        p_v.paragraph_format.space_after = Pt(2)
        r_v = p_v.add_run(v)
        if k == "Lema:":
            r_v.italic = True
            r_v.font.color.rgb = RGBColor(0xb7, 0x79, 0x1f) # Warm gold

        set_cell_background(cell_k, "F7FAFC")
        set_cell_background(cell_v, "F7FAFC")
        set_cell_margins(cell_k, top=60, bottom=60, left=100, right=100)
        set_cell_margins(cell_v, top=60, bottom=60, left=100, right=100)

    doc.add_paragraph() # Spacer

    # Read docs.md content
    with open("docs.md", "r", encoding="utf-8") as f:
        md_lines = f.readlines()

    in_code_block = False
    code_lang = ""
    code_lines = []
    
    in_table = False
    table_rows = []

    def flush_code_block():
        nonlocal code_lines, code_lang
        if not code_lines:
            return
        code_content = "".join(code_lines).strip()
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        cell.width = Inches(6.5)
        set_cell_background(cell, "1E1E1E" if "json" in code_lang or "python" in code_lang or "javascript" in code_lang else "2D3748")
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(code_content)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0) # Light text
        doc.add_paragraph() # Small spacer after code
        code_lines = []
        code_lang = ""

    def flush_table():
        nonlocal table_rows
        if not table_rows:
            return
        cols_count = max(len(r) for r in table_rows)
        t = doc.add_table(rows=len(table_rows), cols=cols_count)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r_idx, r_data in enumerate(table_rows):
            for c_idx, c_text in enumerate(r_data):
                if c_idx < cols_count:
                    cell = t.rows[r_idx].cells[c_idx]
                    set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    if r_idx == 0:
                        set_cell_background(cell, "2B6CB0")
                        run = p.add_run(c_text.strip())
                        run.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    else:
                        set_cell_background(cell, "FFFFFF" if r_idx % 2 == 1 else "EDF2F7")
                        run = p.add_run(c_text.strip())
                        run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
        doc.add_paragraph()
        table_rows = []

    # Skip title metadata from md as already written
    start_parsing = False

    for line in md_lines:
        if line.startswith("## 1. Portada"):
            start_parsing = True

        if not start_parsing:
            continue

        # Handle Code Blocks
        if line.startswith("```"):
            if in_code_block:
                in_code_block = False
                flush_code_block()
            else:
                in_code_block = True
                code_lang = line.strip().replace("```", "").lower()
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Handle Markdown Tables
        if "|" in line and "-|-" in line:
            continue # Header separator
        if "|" in line:
            in_table = True
            cols = [c.strip() for c in line.strip().split("|")[1:-1]]
            if cols:
                table_rows.append(cols)
            continue
        else:
            if in_table:
                in_table = False
                flush_table()

        # Handle Headings
        if line.startswith("## "):
            h = doc.add_heading(level=1)
            run = h.add_run(line.replace("## ", "").strip())
            run.font.name = 'Arial'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1a, 0x36, 0x5d)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            continue
        elif line.startswith("### "):
            h = doc.add_heading(level=2)
            run = h.add_run(line.replace("### ", "").strip())
            run.font.name = 'Arial'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2b, 0x6c, 0xb0)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            continue
        elif line.startswith("#### "):
            h = doc.add_heading(level=3)
            run = h.add_run(line.replace("#### ", "").strip())
            run.font.name = 'Arial'
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2d, 0x37, 0x48)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            continue

        # Handle Bullet Points
        clean_line = line.strip()
        if not clean_line:
            continue

        if clean_line.startswith("* ") or clean_line.startswith("- "):
            p = doc.add_paragraph(style='List Bullet')
            text = clean_line[2:]
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endsWith('**') if hasattr(part, 'endsWith') else (part.startswith('**') and part.endswith('**')):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
            p.paragraph_format.space_after = Pt(3)
            continue

        # Handle Numbered List
        if re.match(r'^\d+\.\s', clean_line):
            p = doc.add_paragraph(style='List Number')
            text = re.sub(r'^\d+\.\s', '', clean_line)
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
            p.paragraph_format.space_after = Pt(3)
            continue

        # Normal Paragraph
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*.*?\*\*)', clean_line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2])
                r.bold = True
            elif part.startswith('*') and part.endswith('*'):
                r = p.add_run(part[1:-1])
                r.italic = True
            elif part.startswith('`') and part.endswith('`'):
                r = p.add_run(part[1:-1])
                r.font.name = 'Consolas'
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(0xC5, 0x30, 0x30)
            else:
                p.add_run(part)

    # Save document
    output_path = "documents/Documentacion_Academia_Lumina_AI.docx"
    doc.save(output_path)
    print(f"Successfully generated DOCX at {output_path}")

if __name__ == "__main__":
    create_document()
